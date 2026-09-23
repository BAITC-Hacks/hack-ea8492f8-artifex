"""Bounded JSON/stdin adapter for the upstream Word loader and optional PDF/Excel parsers."""
import base64
import io
import json
from pathlib import Path
import sys
import tempfile
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from orgsolvency.ingest import read_docx


def extract(name, blob):
    if not blob or len(blob) > 8 * 1024 * 1024:
        raise ValueError("Input must be nonempty and at most 8 MB.")
    suffix = Path(name).suffix.lower()
    notes = []
    if suffix in (".docx", ".xlsx"):
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            if len(archive.infolist()) > 10000 or sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("The expanded document exceeds the safety limit.")
    if suffix == ".docx":
        with tempfile.NamedTemporaryFile(suffix=".docx") as file:
            file.write(blob)
            file.flush()
            text, info = read_docx(file.name)
        if info["autonumbered_unreliable"]:
            notes.append("Word automatic numbering is not reconstructed; affected paragraphs retain [?] anchors.")
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ValueError("PDF input requires the optional pypdf package; install requirements-input.txt.")
        reader = PdfReader(io.BytesIO(blob))
        if reader.is_encrypted or len(reader.pages) > 200:
            raise ValueError("Encrypted PDFs and PDFs over 200 pages are not supported.")
        pages = []
        for i, page in enumerate(reader.pages):
            content = page.extract_text() or ""
            if not content.strip():
                notes.append(f"Page {i + 1} has no readable text; OCR is required and this page is missing from extraction.")
            pages.append(f"[Page {i + 1}]\n{content}")
        text = "\n\n".join(pages)
        if all(not p.split("\n", 1)[-1].strip() for p in pages):
            raise ValueError("This PDF contains no readable text. Supply an OCR text export.")
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValueError("Excel input requires the optional openpyxl package; install requirements-input.txt.")
        workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=False, keep_links=False)
        lines, cells = [], 0
        for sheet in workbook:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows():
                cells += len(row)
                if cells > 100000:
                    raise ValueError("Workbook exceeds 100,000 cells.")
                values = [f"{c.coordinate}: {c.value}" for c in row if c.value is not None]
                if values:
                    lines.append(" | ".join(values))
        workbook.close()
        text = "\n".join(lines)
        notes.append("Spreadsheet formulas are retained as text, never executed; drawings and images are not extracted.")
    elif suffix in (".txt", ".md"):
        text = blob.decode("utf-8-sig")
    else:
        raise ValueError("Supported inputs: .txt, .md, .docx, text PDF, and .xlsx.")
    if not text.strip() or len(text) > 500000:
        raise ValueError("Extracted text must contain 1 to 500,000 characters.")
    return {"text": text, "ingestionNotes": notes}


if __name__ == "__main__":
    try:
        data = json.loads(sys.stdin.read(12 * 1024 * 1024))
        result = extract(data["name"], base64.b64decode(data["base64"], validate=True))
        print(json.dumps(result, ensure_ascii=False))
    except (ValueError, KeyError, zipfile.BadZipFile, UnicodeError) as error:
        print(json.dumps({"error": str(error)}))
        sys.exit(2)
    except Exception:
        print(json.dumps({"error": "Document parsing failed; check the file format."}))
        sys.exit(2)
